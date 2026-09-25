"""Documents of the demo corpus (Э4, Э11): DOCX, PDF, TXT, LOG, TSV.

Numbers are taken from the executed notebooks, so documents and notebooks agree.
Fonts: DejaVu Sans shipped with matplotlib (free licence, Cyrillic support).
"""

from __future__ import annotations

import base64
import io
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

import matplotlib
import nbformat
import numpy as np

FONT_DIR = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
FONT = FONT_DIR / "DejaVuSans.ttf"
FONT_BOLD = FONT_DIR / "DejaVuSans-Bold.ttf"


def notebook_png(nb_path: Path, cell: int) -> bytes:
    """PNG output of a notebook cell (1-based index)."""
    nb = nbformat.read(str(nb_path), as_version=4)
    for out in nb.cells[cell - 1].get("outputs", []):
        data = out.get("data", {})
        if "image/png" in data:
            return base64.b64decode(data["image/png"])
    raise ValueError(f"no image in {nb_path.name} cell {cell}")


# --------------------------------------------------------------------------- DOCX


def _tracked_change(paragraph, before: str, deleted: str, inserted: str, after: str, author: str, date: str) -> None:
    """Append runs with a tracked deletion and insertion (python-docx has no API for revisions)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    paragraph.add_run(before)
    for change_id, (tag, text) in enumerate((("w:del", deleted), ("w:ins", inserted)), start=101):
        change = OxmlElement(tag)
        change.set(qn("w:id"), str(change_id))
        change.set(qn("w:author"), author)
        change.set(qn("w:date"), date)
        run = OxmlElement("w:r")
        t = OxmlElement("w:delText" if tag == "w:del" else "w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        run.append(t)
        change.append(run)
        paragraph._p.append(change)
    paragraph.add_run(after)


def build_report_docx(path: Path, curves_png: bytes) -> None:
    import docx
    from docx.shared import Cm

    doc = docx.Document()
    doc.core_properties.author = "Студент"
    doc.core_properties.title = "Отчёт за март 2026"
    doc.core_properties.created = datetime(2026, 3, 31, 18, 0, tzinfo=timezone.utc)
    doc.sections[0].header.paragraphs[0].text = "Лаборатория машинного обучения — внутренний отчёт"

    doc.add_heading("Отчёт за март 2026", level=0)
    doc.add_heading("1. Цель", level=1)
    doc.add_paragraph(
        "Сравнить линейные модели, градиентный бустинг и многослойный перцептрон на учебных датасетах "
        "scikit-learn и выбрать бейзлайны для дальнейших экспериментов."
    )
    doc.add_heading("2. Результаты", level=1)
    doc.add_paragraph("Метрики посчитаны на валидационной выборке функцией calc_metrics.")
    doc.add_paragraph("Таблица 1 — Метрики моделей на валидации", style="Caption")
    rows = [
        ("Модель", "Датасет", "ROC-AUC", "F1 (macro)"),
        ("LogisticRegression, C=1.0", "breast_cancer", "0.998", "0.9906"),
        ("HistGradientBoosting, lr=0.3", "breast_cancer", "0.9928", "0.9627"),
        ("MLPClassifier + сдвиги", "digits", "0.9998", "0.9781"),
        ("NumpyMLP, cosine", "digits", "0.9947 (тест)", "0.9527 (тест)"),
    ]
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            table.cell(i, j).text = value
    doc.add_paragraph(
        "Лучший ROC-AUC на breast cancer у логистической регрессии; бустинг не дал выигрыша на этой задаче."
    )

    doc.add_heading("2.1. Переобучение", level=2)
    doc.add_paragraph(
        "У ручного MLP на numpy переобучение заметно после 30-й эпохи: train loss продолжает падать "
        "(около 0.07 к эпохе 37), а accuracy на валидации перестаёт расти и держится на уровне 0.9639, "
        "достигнутом на эпохе 27. Ранняя остановка с patience 10 прервала обучение на эпохе 37 и вернула "
        "веса лучшей эпохи."
    )
    p = doc.add_paragraph()
    _tracked_change(p, "Чтобы ослабить переобучение, увеличили weight decay с 1e-4 до ", "5e-4", "1e-3",
                    " и уменьшили learning rate до 0.005.", author="Студент", date="2026-03-30T10:00:00Z")
    doc.add_picture(io.BytesIO(curves_png), width=Cm(15))
    doc.add_paragraph("Рис. 1 — Кривые обучения NumpyMLP: cross-entropy и accuracy на валидации", style="Caption")

    doc.add_heading("3. Выводы и планы", level=1)
    conclusion = doc.add_paragraph()
    run = conclusion.add_run("Аугментация сдвигами на 1 пиксель немного повышает accuracy MLP на digits (0.9722 → 0.9778). ")
    conclusion.add_run("В апреле — калибровка вероятностей и сравнение с CatBoost.")
    from docx.oxml.ns import qn

    comment = doc.add_comment([run], text="Проверьте аугментацию поворотами, а не только сдвигами.",
                              author="А. Петров", initials="АП")
    element = getattr(comment, "_comment_elm", None)
    if element is None:
        element = comment._element
    element.set(qn("w:date"), "2026-04-02T09:30:00Z")  # python-docx stamps "now": keep builds reproducible
    doc.save(str(path))
    normalize_zip(path)


def normalize_zip(path: Path, stamp=(2026, 3, 31, 0, 0, 0)) -> None:
    """Fixed timestamps inside the DOCX archive: rebuilding the corpus gives identical bytes."""
    import zipfile

    with zipfile.ZipFile(path) as zf:
        entries = [(info.filename, zf.read(info.filename)) for info in zf.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in entries:
            zf.writestr(zipfile.ZipInfo(name, date_time=stamp), data, compress_type=zipfile.ZIP_DEFLATED)


# --------------------------------------------------------------------------- PDF


def _fonts():
    from reportlab import rl_config
    from reportlab.pdfbase import pdfmetrics

    rl_config.invariant = 1  # no creation date or random document id: reproducible PDFs
    from reportlab.pdfbase.ttfonts import TTFont

    if "DejaVu" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("DejaVu", str(FONT)))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(FONT_BOLD)))


def _styles():
    from reportlab.lib.styles import ParagraphStyle

    _fonts()
    return {
        "title": ParagraphStyle("title", fontName="DejaVu-Bold", fontSize=15, leading=19, spaceAfter=6),
        "author": ParagraphStyle("author", fontName="DejaVu", fontSize=9, leading=12, spaceAfter=8),
        "h1": ParagraphStyle("h1", fontName="DejaVu-Bold", fontSize=11, leading=14, spaceBefore=8, spaceAfter=4),
        "body": ParagraphStyle("body", fontName="DejaVu", fontSize=9, leading=12, spaceAfter=5),
        "caption": ParagraphStyle("caption", fontName="DejaVu", fontSize=8, leading=10, spaceBefore=3, spaceAfter=8),
    }


def build_paper_pdf(path: Path, confusion_png: bytes) -> None:
    """Two-column short paper: reading order across columns, a table and a figure."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.platypus import BaseDocTemplate, Frame, FrameBreak, Image, PageTemplate, Paragraph, Table, TableStyle

    st = _styles()
    width, height = A4
    margin, gap, top = 1.8 * cm, 0.8 * cm, 5.2 * cm
    col_w = (width - 2 * margin - gap) / 2
    first = [
        Frame(margin, height - margin - top, width - 2 * margin, top, id="head"),
        Frame(margin, margin, col_w, height - 2 * margin - top, id="c1"),
        Frame(margin + col_w + gap, margin, col_w, height - 2 * margin - top, id="c2"),
    ]
    later = [Frame(margin, margin, col_w, height - 2 * margin, id="l1"),
             Frame(margin + col_w + gap, margin, col_w, height - 2 * margin, id="l2")]
    doc = BaseDocTemplate(str(path), pagesize=A4, title="Shift Augmentation for Small-Image Digit Classification",
                          author="Demo corpus")
    doc.addPageTemplates([PageTemplate(id="first", frames=first, autoNextPageTemplate="later"),
                          PageTemplate(id="later", frames=later)])
    body = lambda t: Paragraph(t, st["body"])  # noqa: E731
    story = [
        Paragraph("Shift Augmentation for Small-Image Digit Classification", st["title"]),
        Paragraph("Research archive note, March 2026", st["author"]),
        body("<b>Abstract.</b> We study whether one-pixel shift augmentation helps a multilayer perceptron on the "
             "8x8 digits dataset. With identical hyperparameters, augmentation raises validation accuracy from "
             "0.9722 to 0.9778 and macro F1 from 0.9722 to 0.9781."),
        FrameBreak(),
        Paragraph("1 Introduction", st["h1"]),
        body("Small images leave little room for invariances to be learned from data. Classic remedies are "
             "translations, rotations and elastic distortions. Here we isolate the effect of translations on a "
             "fully connected network, which, unlike a convolutional one, has no built-in shift invariance."),
        Paragraph("2 Method", st["h1"]),
        body("Every training image is copied four times and shifted by one pixel left, right, up and down, with "
             "zero padding. The training set therefore grows five times, from 1077 to 5385 images. Validation and "
             "test images are not augmented. The split is stratified 60/20/20 with seed 0."),
        body("The network has one hidden layer of 128 units, learning rate 3e-3, L2 penalty 1e-4, batch size 64 "
             "and at most 200 iterations; inputs are standardised."),
        Paragraph("3 Experiments", st["h1"]),
        Paragraph("Table 1: Validation metrics with and without shift augmentation.", st["caption"]),
    ]
    table = Table([["Training data", "Accuracy", "Macro F1", "ROC-AUC"],
                   ["Original", "0.9722", "0.9722", "0.9983"],
                   ["With shifts", "0.9778", "0.9781", "0.9998"]], colWidths=[2.6 * cm, 1.6 * cm, 1.6 * cm, 1.6 * cm])
    table.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), "DejaVu"), ("FONTSIZE", (0, 0), (-1, -1), 8),
                               ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                               ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke)]))
    story += [
        table,
        body("The gain is small but consistent across metrics. Most remaining errors are confusions between "
             "visually similar digits, as the confusion matrix in Figure 1 shows."),
        Image(io.BytesIO(confusion_png), width=col_w, height=col_w),
        Paragraph("Figure 1: Confusion matrix of the augmented MLP on the validation set.", st["caption"]),
        Paragraph("4 Conclusion", st["h1"]),
        body("One-pixel shifts are a cheap augmentation that slightly improves an MLP on 8x8 digits. Rotations "
             "and a convolutional baseline are left for future work."),
    ]
    doc.build(story)


def build_lecture_pdf(path: Path) -> None:
    """Single-column Russian lecture notes with a table."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

    st = _styles()
    doc = SimpleDocTemplate(str(path), pagesize=A4, title="Конспект лекции: регуляризация", author="Demo corpus",
                            leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm)
    body = lambda t: Paragraph(t, st["body"])  # noqa: E731
    story = [
        Paragraph("Конспект лекции 6. Регуляризация линейных моделей", st["title"]),
        Paragraph("Курс «Машинное обучение», весенний семестр", st["author"]),
        Paragraph("1. Зачем нужна регуляризация", st["h1"]),
        body("Регуляризация добавляет к функции потерь штраф за сложность модели и уменьшает дисперсию оценок. "
             "Без неё линейная модель на данных с коррелированными признаками получает большие по модулю веса "
             "противоположных знаков и плохо обобщается."),
        Paragraph("2. L1 и L2", st["h1"]),
        body("L2-регуляризация (ridge) штрафует сумму квадратов весов. Она равномерно сжимает все веса к нулю, но "
             "почти никогда не делает их точно нулевыми и устойчива к мультиколлинеарности."),
        body("L1-регуляризация (lasso) штрафует сумму модулей весов. Из-за угловой геометрии допустимой области "
             "решение часто лежит на осях, и часть весов становится ровно нулевой — L1 выполняет отбор признаков."),
        Paragraph("Таблица 1. Сравнение L1 и L2", st["caption"]),
    ]
    table = Table([["", "L1 (lasso)", "L2 (ridge)"],
                   ["Штраф", "сумма |w|", "сумма w²"],
                   ["Нулевые веса", "да, разреженность", "нет"],
                   ["Коррелированные признаки", "выбирает один", "делит вес"],
                   ["Решение", "итеративное", "аналитическое"]], colWidths=[5 * cm, 4.5 * cm, 4.5 * cm])
    table.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), "DejaVu"), ("FONTSIZE", (0, 0), (-1, -1), 9),
                               ("GRID", (0, 0), (-1, -1), 0.4, colors.grey)]))
    story += [
        table,
        Paragraph("3. Weight decay в нейросетях", st["h1"]),
        body("В градиентном спуске L2-штраф эквивалентен weight decay: на каждом шаге веса умножаются на "
             "(1 − η·λ). Для SGD это одно и то же, для адаптивных оптимизаторов вроде Adam — нет, поэтому "
             "используют AdamW с отделённым weight decay."),
        Paragraph("4. Как выбирать силу регуляризации", st["h1"]),
        body("Коэффициент подбирают по валидации или кросс-валидации на логарифмической сетке, например "
             "0.01, 0.1, 1, 10, 100. Признаки перед этим стандартизируют, иначе штраф действует неравномерно."),
    ]
    doc.build(story)


def build_scanned_pdf(path: Path) -> None:
    """A 'scanned' page: text rendered to a slightly rotated noisy image, no text layer."""
    _fonts()
    from PIL import Image, ImageDraw, ImageFilter, ImageFont
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    lines = [
        ("Протокол встречи научной группы", 44),
        ("20 марта 2026 г.", 32),
        ("", 32),
        ("1. Обсудили результаты экспериментов на digits и breast cancer.", 32),
        ("2. Срок сдачи главы 2 диссертации — 10 апреля.", 32),
        ("3. Подготовить демо системы поиска по архиву к 25 апреля.", 32),
        ("4. Добавить в отчёт калибровку вероятностей (ECE).", 32),
        ("", 32),
        ("Следующая встреча — 3 апреля в 15:00.", 32),
    ]
    img = Image.new("L", (1654, 2339), 250)  # A4 at 200 dpi
    draw = ImageDraw.Draw(img)
    y = 180
    for text, size in lines:
        font = ImageFont.truetype(str(FONT_BOLD if size > 40 else FONT), size)
        draw.text((150, y), text, fill=25, font=font)
        y += int(size * 1.9)
    img = img.rotate(0.6, resample=Image.BICUBIC, fillcolor=250).filter(ImageFilter.GaussianBlur(0.6))
    noise = np.random.default_rng(0).normal(0, 6, (img.height, img.width))
    img = Image.fromarray(np.clip(np.asarray(img, dtype=float) + noise, 0, 255).astype("uint8"))
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=70)  # scanners produce JPEG; PNG of noise is huge
    buf.seek(0)
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setTitle("scan")
    c.drawImage(ImageReader(buf), 0, 0, width=A4[0], height=A4[1])
    c.showPage()
    c.save()


# --------------------------------------------------------------------------- TXT, LOG, TSV


NOTES = """Идеи на апрель

- Попробовать аугментацию поворотами на digits (углы ±10°), сравнить со сдвигами.
- CatBoost на breast cancer вместо HistGradientBoosting.
- Калибровка вероятностей логистической регрессии: посчитать ECE, попробовать Platt scaling.
- Разобраться, почему step-расписание в numpy-MLP хуже constant.

Дедлайн отчёта за апрель — 15 апреля.
"""


def build_notes_txt(path: Path) -> None:
    path.write_bytes(NOTES.replace("\n", "\r\n").encode("cp1251"))  # an old Windows editor


def build_training_log(path: Path, corpus_root: Path) -> None:
    """Real training run of NumpyMLP with the parameters of 08_numpy_mlp, logged per epoch."""
    sys.dont_write_bytecode = True  # no __pycache__ inside the corpus
    sys.path.insert(0, str(corpus_root))
    try:
        from ml_toolkit.data import make_split
        from ml_toolkit.features import make_preprocessor
        from ml_toolkit.trainer import NumpyMLP, Trainer
    finally:
        sys.path.pop(0)
    split = make_split("digits", seed=0)
    pre = make_preprocessor().fit(split.X_train)
    X_tr, X_va = pre.transform(split.X_train), pre.transform(split.X_val)
    model = NumpyMLP(64, 128, 10, seed=0)
    trainer = Trainer(model, lr=0.005, epochs=80, batch_size=64, schedule="cosine", momentum=0.9,
                      weight_decay=1e-3, patience=10, seed=0)
    history = trainer.fit(X_tr, split.y_train, X_va, split.y_val)
    t0 = datetime(2026, 3, 14, 10, 15, 2)
    out = [f"{t0:%Y-%m-%d %H:%M:%S} INFO start training NumpyMLP(64-128-10) lr=0.005 batch=64 schedule=cosine "
           f"weight_decay=0.001 patience=10"]
    for e in range(len(history.train_loss)):
        ts = t0 + timedelta(seconds=3 * (e + 1))
        out.append(f"{ts:%Y-%m-%d %H:%M:%S} INFO epoch={e} lr={history.lr[e]:.5f} "
                   f"train_loss={history.train_loss[e]:.4f} val_loss={history.val_loss[e]:.4f} "
                   f"val_acc={history.val_accuracy[e]:.4f}")
    ts = t0 + timedelta(seconds=3 * (len(history.train_loss) + 1))
    out.append(f"{ts:%Y-%m-%d %H:%M:%S} INFO early stopping at epoch {trainer.stopped_epoch}; "
               f"restored best epoch {history.best_epoch()} (val_acc={max(history.val_accuracy):.4f})")
    path.write_text("\n".join(out) + "\n", encoding="utf-8", newline="\n")


def build_dataset_stats(path: Path) -> None:
    from sklearn import datasets

    rows = ["dataset\tn_samples\tn_features\tn_classes\ttask"]
    for name, loader, task in [("breast_cancer", datasets.load_breast_cancer, "classification"),
                               ("wine", datasets.load_wine, "classification"),
                               ("digits", datasets.load_digits, "classification"),
                               ("diabetes", datasets.load_diabetes, "regression")]:
        b = loader()
        n_classes = len(np.unique(b.target)) if task == "classification" else "-"
        rows.append(f"{name}\t{b.data.shape[0]}\t{b.data.shape[1]}\t{n_classes}\t{task}")
    path.write_text("\n".join(rows) + "\n", encoding="utf-8", newline="\n")


def build_documents(out: Path) -> None:
    for sub in ("docs", "papers", "notes", "logs", "exports"):
        (out / sub).mkdir(parents=True, exist_ok=True)
    curves = notebook_png(out / "experiments" / "08_numpy_mlp.ipynb", 6)
    confusion = notebook_png(out / "experiments" / "04_digits_augmentation.ipynb", 9)
    build_report_docx(out / "docs" / "report_march.docx", curves)
    build_paper_pdf(out / "papers" / "shift_augmentation_note.pdf", confusion)
    build_lecture_pdf(out / "papers" / "lecture_regularization.pdf")
    build_scanned_pdf(out / "papers" / "meeting_notes_scan.pdf")
    build_notes_txt(out / "notes" / "ideas_april.txt")
    build_training_log(out / "logs" / "train_numpy_mlp.log", out)
    build_dataset_stats(out / "exports" / "dataset_stats.txt")
