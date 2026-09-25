"""Э4 / Э11: DOCX, TXT (markdown, prose, log, table) parsers; PDF text-layer check."""

from pathlib import Path

import pytest

from rag_agent.config import ChunkingConfig
from rag_agent.ingest import iter_corpus, parse_file
from rag_agent.ingest.txt_parser import classify, decode, texttiling
from rag_agent.schema import EdgeType, NodeType

docx = pytest.importorskip("docx")


def parsed(root: Path, name: str, **cfg):
    cf = next(f for f in iter_corpus(root, [Path(name).suffix]) if f.rel_path == name)
    return parse_file(cf, ChunkingConfig(**cfg))


def make_docx(path: Path) -> None:
    import sys

    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
    try:
        from demo_documents import _tracked_change
    finally:
        sys.path.pop(0)
    d = docx.Document()
    d.core_properties.author = "Tester"
    d.add_heading("Report", level=0)
    d.add_heading("Results", level=1)
    d.add_paragraph("Accuracy was measured on the validation split.")
    d.add_paragraph("Table 1 — Scores", style="Caption")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "model", "score"
    table.cell(1, 0).text, table.cell(1, 1).text = "logreg", "0.998"
    d.add_heading("Tuning", level=2)
    p = d.add_paragraph()
    _tracked_change(p, "The learning rate was set to ", "0.1", "0.01", " after tuning.", author="A", date="2026-01-01T00:00:00Z")
    note = d.add_paragraph()
    run = note.add_run("Early stopping patience is 10.")
    d.add_comment([run], text="Why not 5?", author="Reviewer")
    d.save(str(path))


def test_docx_structure_revisions_comments(tmp_path: Path):
    make_docx(tmp_path / "r.docx")
    result = parsed(tmp_path, "r.docx")
    nodes = result.nodes
    assert nodes[0].metadata["author"] == "Tester" and nodes[0].metadata["n_revisions"] == 1
    sections = [n for n in nodes if n.node_type == NodeType.SECTION]
    tuning = next(n for n in sections if n.location.section == "Report > Results > Tuning")
    assert "set to 0.01 after tuning" in tuning.text and "0.1 " not in tuning.text  # accepted changes only
    # the title-only heading travels into the first real chunk
    assert sections[0].text.startswith("Report\nResults")
    table = next(n for n in nodes if n.node_type == NodeType.TABLE)
    assert table.title == "Table 1 — Scores" and "| logreg | 0.998 |" in table.text
    rev = next(n for n in nodes if n.node_type == NodeType.REVISION)
    assert "Удалено: «0.1»" in rev.text and "Вставлено: «0.01»" in rev.text and rev.metadata["author"] == "A"
    comment = next(n for n in nodes if n.node_type == NodeType.COMMENT)
    assert "Reviewer" in comment.text and "Why not 5?" in comment.text and "patience is 10" in comment.context
    assert tuning.embedding_text(True).startswith("File: r.docx | Tuning | Report > Results > Tuning")


def test_docx_plain_baseline_loses_structure(tmp_path: Path):
    make_docx(tmp_path / "r.docx")
    nodes = parsed(tmp_path, "r.docx", docx="plain").nodes[1:]
    assert all(n.node_type == NodeType.SECTION and n.location.section is None for n in nodes)
    assert not any(n.node_type in (NodeType.REVISION, NodeType.COMMENT, NodeType.TABLE) for n in nodes)
    assert "set to  after tuning" in nodes[0].text  # naive extraction drops the tracked insertion


def test_decode_cp1251_and_classify(tmp_path: Path):
    text, enc = decode("Дедлайн отчёта — 15 апреля. Заметки о переобучении и аугментации.".encode("cp1251"))
    assert "Дедлайн" in text and enc.replace("_", "").lower() in ("cp1251", "windows1251")
    assert classify("2026-01-01 10:00:00 epoch=1 loss=0.5\n2026-01-01 10:00:01 epoch=2 loss=0.4\n") == "log"
    assert classify("a\tb\tc\n1\t2\t3\n4\t5\t6\n") == "table"
    assert classify("# Title\n\nSome text.\n\n## Part\nMore.") == "markdown"
    assert classify("import os\ndef f():\n    return 1\n") == "code"
    assert classify("Just a normal sentence. Another one follows here.") == "prose"


def test_log_metrics_summary(tmp_path: Path):
    lines = [f"2026-03-14 10:00:{e:02d} INFO epoch={e} train_loss={1 / (e + 1):.3f} val_acc={0.5 + 0.1 * min(e, 3):.2f}"
             for e in range(6)]
    (tmp_path / "train.log").write_text("\n".join(lines), encoding="utf-8")
    result = parsed(tmp_path, "train.log")
    summary = next(n for n in result.nodes if n.node_type == NodeType.SUMMARY)
    assert "val_acc: максимум 0.8 (epoch 3)" in summary.text
    assert "train_loss: минимум 0.167 (epoch 5)" in summary.text
    assert result.nodes[0].metadata["metrics"]["val_acc"]["best_step"] == 3
    assert any(n.node_type == NodeType.LOG_CHUNK for n in result.nodes)


def test_table_export_and_markdown_sections(tmp_path: Path):
    (tmp_path / "stats.txt").write_text("dataset\tn\nwine\t178\ndigits\t1797\n", encoding="utf-8")
    table = parsed(tmp_path, "stats.txt").nodes[1]
    assert table.node_type == NodeType.TABLE and "| wine | 178 |" in table.text
    (tmp_path / "notes.md").write_text("# Plan\n\nIntro.\n\n## April\n\nTry rotations.\n", encoding="utf-8")
    sections = parsed(tmp_path, "notes.md").nodes[1:]
    assert [n.location.section for n in sections] == ["Plan", "Plan > April"]
    assert sections[1].location.line_start == 5


def test_texttiling_splits_topics():
    a = "Cats purr and sleep. Cats hunt mice at night. A cat likes warm places. Kittens play with cats. " * 2
    b = "Stocks fell sharply today. Markets and stocks react to rates. Investors sold stocks. Bond markets rose. " * 2
    pieces = texttiling(a + b, block=2)
    assert len(pieces) >= 2 and "Cats" in pieces[0] and "Stocks" in pieces[-1]


def test_pdf_text_layer_check(tmp_path: Path):
    pytest.importorskip("pypdfium2")
    reportlab = pytest.importorskip("reportlab")  # noqa: F841
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(tmp_path / "t.pdf"))
    c.drawString(72, 720, "A page with a real text layer, long enough to count.")
    c.showPage()
    c.showPage()  # empty page: looks scanned
    c.save()
    from rag_agent.ingest.pdf_parser import pages_without_text

    assert pages_without_text(tmp_path / "t.pdf", min_chars=30) == (2, [2])


def test_docx_edges_link_figures(tmp_path: Path):
    from docx.shared import Cm

    Image = pytest.importorskip("PIL.Image")

    Image.new("RGB", (10, 10), "white").save(tmp_path / "i.png")
    d = docx.Document()
    d.add_heading("Plots", level=1)
    d.add_paragraph("The curve is shown below.")
    d.add_picture(str(tmp_path / "i.png"), width=Cm(2))
    d.add_paragraph("Figure 1 — Learning curve", style="Caption")
    d.save(str(tmp_path / "f.docx"))
    result = parsed(tmp_path, "f.docx")
    fig = next(n for n in result.nodes if n.node_type == NodeType.FIGURE)
    assert fig.text == "Figure 1 — Learning curve" and fig.embed
    assert any(e.type == EdgeType.ILLUSTRATES and e.src == fig.id for e in result.edges)
